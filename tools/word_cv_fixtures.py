"""Synthetic Word regression inputs and independent OOXML readback; never open user documents."""
import json
import sys
from pathlib import Path
from docx import Document

action, folder = sys.argv[1:3]
root = Path(folder).resolve()
if action == 'prepare':
    root.mkdir(exist_ok=False, parents=True)
    (root / '.h2-agent-test-fixture').write_text('Synthetic Word CV acceptance only', encoding='utf-8')
    for n in range(1, 7):
        doc = Document()
        doc.add_paragraph().add_run(f'H2 CV TEST {n}').bold = True
        body = doc.add_paragraph()
        body.add_run('' if n in (2, 5) else 'Nội dung mẫu cần thay').italic = n == 3
        doc.add_paragraph().add_run(f'KEEP-{n}').italic = True
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = f'TABLE-{n}'
        table.cell(0, 1).text = '42'
        doc.sections[0].header.paragraphs[0].text = f'HEADER-{n}'
        doc.sections[0].footer.paragraphs[0].text = f'FOOTER-{n}'
        doc.save(root / f'h2-cv-regression-{n}.docx')
    print(root)
elif action == 'verify':
    if not (root / '.h2-agent-test-fixture').is_file():
        raise RuntimeError('Missing synthetic marker')
    results = []
    for n in range(1, 7):
        name = f'cv-script-result-{n}.docx' if n <= 3 else f'cv-gemma-result-{n-3}.docx'
        path = root / name
        if not path.exists():
            results.append({'file': name, 'checks': {'exists': False}})
            continue
        source = Document(root / f'h2-cv-regression-{n}.docx')
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        checks = {
            'cv_paragraph_count': len(texts) >= 7,
            'title_bold': all(r.bold for r in doc.paragraphs[0].runs if r.text),
            'keep_once': texts.count(f'KEEP-{n}') == 1,
            'keep_format': all(r.italic for p in doc.paragraphs if p.text == f'KEEP-{n}' for r in p.runs if r.text),
            'table': [[c.text for c in r.cells] for r in doc.tables[0].rows] == [[f'TABLE-{n}', '42']],
            'header': doc.sections[0].header.paragraphs[0].text == f'HEADER-{n}',
            'footer': doc.sections[0].footer.paragraphs[0].text == f'FOOTER-{n}',
            'page_setup': all(getattr(doc.sections[0], k) == getattr(source.sections[0], k) for k in ('page_width', 'page_height', 'top_margin', 'left_margin', 'bottom_margin', 'right_margin')),
            'vietnamese': any('Giới' in t or 'giới' in t or 'Kỹ' in t or 'kỹ' in t for t in texts),
        }
        results.append({'file': name, 'checks': checks, 'paragraphs': texts})
    report = {'passed': all(all(r['checks'].values()) for r in results), 'results': results}
    (root / 'independent-readback.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    sys.exit(0 if report['passed'] else 1)
else:
    raise ValueError(action)
