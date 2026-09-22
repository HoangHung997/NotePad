"""Small synthetic fixtures and independent readback for the opt-in production Agent probe."""
import argparse
import hashlib
import importlib.util
import json
from pathlib import Path


def prepare(root):
    from docx import Document
    from openpyxl import Workbook
    root.mkdir(parents=True, exist_ok=False)
    spec = importlib.util.spec_from_file_location("ocr_fixture", Path(__file__).parent / "ocr/smoke.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod.fixtures(root / "ocr-fixtures")
    create, edit, delete, web = [], [], [], []
    for i in range(1, 4):
        work = root / f"round-{i}"
        work.mkdir()
        (work / ".h2-agent-test-fixture").write_text("Synthetic data only. Test deletion permitted inside this directory.", encoding="utf-8")
        (work / "source.txt").write_text(f"H2-TEST-{i}\nTên dự án: Kiểm thử Agent\nAlpha: 12\nBeta: 30\nTổng: 42\nGiữ nguyên mã KEEP-{i}\n", encoding="utf-8")
        document = Document()
        document.add_heading(f"H2-TEST-{i}", 0)
        document.add_paragraph("Bản dự thảo chưa sửa.")
        document.add_paragraph(f"KEEP-{i}")
        document.save(work / "seed.docx")
        book = Workbook()
        sheet = book.active; sheet.title = "Data"
        sheet.append(["Item", "Quantity"]); sheet.append(["Alpha", 12]); sheet.append(["Beta", 30]); sheet.append(["Total", "=SUM(B2:B3)"])
        sheet["D1"] = f"KEEP-{i}"; sheet["A1"].font = __import__("openpyxl").styles.Font(bold=True)
        book.save(work / "seed.xlsx")
        create.append(dict(id=f"create-{i}", workspace=str(work), prompt=f"Đọc source.txt. Dùng run_python và publish_artifact tạo report.docx có tiêu đề H2-TEST-{i}, đoạn tiếng Việt báo cáo Alpha 12, Beta 30, tổng 42 và mã KEEP-{i}; tiêu đề in đậm. Tạo report.xlsx sheet Data, A1 Item B1 Quantity, A2 Alpha B2 12, A3 Beta B3 30, A4 Total B4 công thức =SUM(B2:B3), D1 KEEP-{i}; hàng tiêu đề đậm. Tạo report.pdf có chữ H2-TEST-{i}, Alpha 12, Beta 30, Total 42, KEEP-{i}. Công bố đủ ba file vào thư mục làm việc đúng tên; kiểm tra đọc lại tất cả."))
        edit.append(dict(id=f"edit-{i}", workspace=str(work), prompt=f"Sửa đúng 3 tệp report.docx report.xlsx report.pdf trong workspace. Word: thêm một đoạn EDITED-{i}, giữ các đoạn cũ và định dạng tiêu đề. Excel: B2 thành 15, B4 vẫn =SUM(B2:B3), D1 và định dạng giữ nguyên; tổng mới 45. PDF: thêm trang cuối ghi EDITED-{i} Total 45, giữ trang cũ. Có thể dùng run_python trên bản sao rồi publish_artifact với hash hiện tại để thay thế. Đọc lại kiểm tra cả 3 bản đã lưu."))
        delete.append(dict(id=f"delete-{i}", workspace=str(work), fullAccess=True, prompt="Tạo thư mục delete-only trong workspace, sao chép report.docx, report.xlsx, report.pdf và source.txt vào đó; sửa tên các bản sao có tiền tố disposable-. Kiểm tra tồn tại rồi xóa đúng 4 bản sao disposable đó. Trước khi xóa kiểm tra đường dẫn tuyệt đối nằm trong delete-only. Giữ nguyên toàn bộ file gốc. Ghi deletion-log.json trong workspace ghi tên 4 tệp, existed_before và absent_after cho từng tệp; kiểm tra thư mục delete-only trống. Dùng công cụ thực thi PowerShell khi cần."))
        web.append(dict(id=f"news-{i}", workspace=str(work), prompt="Truy cập trực tiếp https://vnexpress.net/rss/tin-moi-nhat.rss và https://feeds.bbci.co.uk/news/world/rss.xml bằng công cụ web. Đọc tin thật, chọn tối đa 5 tin liên quan công nghệ/kinh tế/quốc tế, loại trùng tiêu đề hoặc URL; chỉ giữ tin có ngày xuất bản trong 72 giờ gần nhất tính theo đồng hồ máy. Nếu nguồn lỗi ghi rõ lỗi; không tự tạo tin. Lưu news.json gồm fetched_at_utc, sources, articles (title, url, published, category, summary_vi). Kèm một file news.md tóm tắt tiếng Việt có link dẫn. Nếu không có tin đạt điều kiện hãy trả danh sách rỗng kèm giải thích. Đọc lại file để kiểm tra."))
    for name, cases in [("create", create), ("edit", edit), ("delete", delete), ("news", web)]:
        (root / f"{name}.json").write_text(json.dumps(cases, ensure_ascii=False, indent=2), encoding="utf-8")
    print(root)


def verify(root, stage):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    checks = []
    for i in range(1, 4):
        work = root / f"round-{i}"
        def check(name, value):
            checks.append(dict(round=i, check=name, passed=bool(value)))
        try:
            doc = Document(work / "report.docx")
            text = "\n".join(p.text for p in doc.paragraphs)
            check("word-title-preserved", f"H2-TEST-{i}" in text)
            check("word-body-preserved", f"KEEP-{i}" in text and "Alpha" in text and "Beta" in text)
            book = load_workbook(work / "report.xlsx")
            sheet = book["Data"]
            check("excel-quantity", sheet["B2"].value == (15 if stage != "create" else 12) and sheet["B3"].value == 30)
            check("excel-formula", sheet["B4"].value == "=SUM(B2:B3)")
            check("excel-preservation", sheet["D1"].value == f"KEEP-{i}" and sheet["A1"].font.bold)
            pdf = PdfReader(work / "report.pdf")
            check("pdf-original-page", f"H2-TEST-{i}" in pdf.pages[0].extract_text() and f"KEEP-{i}" in pdf.pages[0].extract_text())
            if stage != "create":
                check("word-edited", f"EDITED-{i}" in text)
                check("pdf-edited", len(pdf.pages) >= 2 and f"EDITED-{i}" in pdf.pages[-1].extract_text() and "45" in pdf.pages[-1].extract_text())
            if stage == "delete":
                log = json.loads((work / "deletion-log.json").read_text(encoding="utf-8-sig"))
                check("delete-only-empty", not list((work / "delete-only").iterdir()))
                check("delete-originals-preserved", all((work / name).exists() for name in ["report.docx", "report.xlsx", "report.pdf", "source.txt"]))
                check("delete-log", bool(log))
        except Exception as e:
            checks.append(dict(round=i, check="readback", passed=False, error=str(e)))
    result = dict(stage=stage, checks=checks, passed=all(c["passed"] for c in checks))
    (root / f"verify-{stage}.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=["prepare", "create", "edit", "delete"])
    parser.add_argument("root", type=Path)
    args = parser.parse_args()
    if args.action == "prepare": prepare(args.root.resolve())
    else: raise SystemExit(verify(args.root.resolve(), args.action))
